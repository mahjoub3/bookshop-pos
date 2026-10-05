"""Excel exports with openpyxl."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def table_to_xlsx(*, sheet_title, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
    for row in rows:
        ws.append(list(row))
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(width, 50)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
